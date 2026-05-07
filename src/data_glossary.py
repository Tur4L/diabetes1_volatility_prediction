import pandas as pd
from pathlib import Path
import argparse, os, sys
import numpy as np
from docx import Document

def read_data(file_path: Path) -> pd.DataFrame:
    ext = file_path.suffix.lower()
    if ext == '.csv':
        return pd.read_csv(file_path)
    elif ext == '.txt':
        return pd.read_csv(file_path, sep="|")
    else:
        raise ValueError(f'Unsupported file type: {ext}')
    
def safe_min(series: pd.Series):
    try:
        if pd.api.types.is_numeric_dtype(series):
            return np.nanmin(series.to_numpy(dtype=float))
        if pd.api.types.is_datetime64_any_dtype(series):
            return pd.to_datetime(series, errors="coerce").min()
        return np.nan
    except Exception:
        return np.nan

def safe_max(series: pd.Series):
    try:
        if pd.api.types.is_numeric_dtype(series):
            return np.nanmax(series.to_numpy(dtype=float))
        if pd.api.types.is_datetime64_any_dtype(series):
            return pd.to_datetime(series, errors="coerce").max()
        return np.nan
    except Exception:
        return np.nan
    
def summarize_column(s: pd.Series, max_values: int=20):
    if pd.api.types.is_numeric_dtype(s):
        dtype = "numeric"
        s2 = s
    elif pd.api.types.is_datetime64_any_dtype(s):
        dtype = "datetime"
        s2 = s
    elif pd.api.types.is_object_dtype(s):
        dtype = "categorical"
        s2 = s
    else: 
        raise ValueError(f"Mixed dtypes in series {s}: {s.dtype}")
    
    col_min = safe_min(s)
    col_max = safe_max(s)

    possible_values = ""
    if dtype == 'categorical':
        vc = s2.fillna("NA").value_counts().head(max_values)
        vals = [str(v) for v in vc.index.tolist()]
        if s2.nunique(dropna=False) > max_values:
            vals.append("...")
        possible_values = ", ".join(vals)

    return dtype, col_min, col_max, possible_values

def to_docx(report_path: Path, study_name: str, tables_meta: list, per_table_dicts: dict):
    doc = Document()
    doc.add_heading('Data Glossary', level=0)
    if study_name:
        doc.add_paragraph().add_run(f"Study: {study_name}").bold = True

    doc.add_heading('Data Table List:', level=1)
    for item in tables_meta:
        line = item["table_name"] + (f" — {item.get('description','')}" if item.get("description") else "")
        doc.add_paragraph(line, style='List Bullet')

    doc.add_heading('Data Table Details:', level=1)
    for item in tables_meta:
        tname = item["table_name"]
        doc.add_heading(tname, level=2)
        if item.get("description"):
            doc.add_paragraph(f"Description: {item['description']}")

        table = doc.add_table(rows=1, cols=6)
        table.autofit = False
        table.allow_autofit = False

        section = doc.sections[0]
        page_width_emu = section.page_width - section.left_margin - section.right_margin  # EMUs (int)

        fractions = [0.25, 0.25, 0.08, 0.08, 0.10, 0.24]  # must sum ~1.0
        for i, f in enumerate(fractions):
            table.columns[i].width = int(page_width_emu * f)  # <-- cast to int EMUs
            
        for i, h in enumerate(["Field_Name","Description","Min","Max","Units","Possible_Values"]):
            table.rows[0].cells[i].text = h

        for row in per_table_dicts[tname]:
            cells = table.add_row().cells
            cells[0].text = str(row.get("Field_Name"))
            cells[1].text = str(row.get("Description","") or "")
            cells[2].text = "" if pd.isna(row.get("Min")) else str(row.get("Min"))
            cells[3].text = "" if pd.isna(row.get("Max")) else str(row.get("Max"))
            cells[4].text = str(row.get("Units","") or "")
            cells[5].text = str(row.get("Possible_Values","") or "")
        doc.add_paragraph()
    doc.save(report_path)
    return True, None

def main():
    ap = argparse.ArgumentParser(description="Create a per-study Data Glossary (DOCX).")
    ap.add_argument("--input-folder", required=True, help="Path to study folder containing datasets.")
    ap.add_argument("--output-prefix", required=True, help="Prefix for output files (without extension).")
    ap.add_argument("--study-name", default="", help="Optional study name.")
    ap.add_argument("--desc-map", default="", help="Optional CSV/Excel with columns: field_name, description")
    ap.add_argument("--units-map", default="", help="Optional CSV/Excel with columns: field_name, units")
    ap.add_argument("--max-values", type=int, default=20, help="Max unique values shown under Possible_Values.")
    args = ap.parse_args()

    in_dir = Path(args.input_folder)
    if not in_dir.exists() or not in_dir.is_dir():
        print(f"[ERROR] Not a directory: {in_dir}", file=sys.stderr); sys.exit(2)

    files = []
    for root, _, fnames in os.walk(in_dir):
        for fn in fnames:
            if Path(fn).suffix.lower() in ['.csv', '.txt']:
                files.append(Path(root) / fn)
    if not files:
        print("[WARN] No supported data files found.", file=sys.stderr)

    tables_meta, per_table_dicts = [], {}

    for fp in sorted(files):
        tname = fp.stem
        try:
            df = read_data(fp)
        except Exception as e:
            print(f"Skipping {fp.name}: {e}", file=sys.stderr)
            continue

        rows = []
        for col in df.columns:
            dtype, col_min, col_max, poss = summarize_column(df[col], max_values= args.max_values)
            nm = str(col).strip().lower()
            units = ""
            if "mg/dl" in nm or "mgdl" in nm: units = "mg/dL"
            elif "mmol/l" in nm or "mmol" in nm: units = "mmol/L"
            elif nm.endswith("_kg") or "weight" in nm: units = units or "kg"
            elif nm.endswith("_cm") or "height" in nm: units = units or "cm"

            rows.append({"Field_Name": col, "Min": col_min, "Max": col_max,
                         "Units": units, "Possible_Values": poss})

        out_df = pd.DataFrame(rows, columns=["Field_Name","Description","Min","Max","Units","Possible_Values"])

        tables_meta.append({"table_name": tname, "description": ""})
        per_table_dicts[tname] = rows

    docx_path = Path(f"{args.output_prefix}_glossary.docx")
    ok, err = to_docx(docx_path, args.study_name, tables_meta, per_table_dicts)
    pd.DataFrame(tables_meta).to_csv(Path(f"{args.output_prefix}_data_table_list.csv"), index=False)
    print("[DONE] Glossary created:")
    print(f"- {docx_path if ok else str(Path(f'{args.output_prefix}_glossary.md'))}")
    print(f"- {args.output_prefix}_data_table_list.csv")

if __name__ == "__main__":
    main()

#python3 src/data_glossary.py --input-folder ./data/studies/jaeb_healthy/original --output-prefix ./data/studies/jaeb_healthy/original --study-name "JAEB Healthy" --desc-map "" --units-map "" --max-values 3
